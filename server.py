from fastapi import FastAPI, Form, Request, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse
from typing import List
from pathlib import Path
from datetime import datetime
from contextlib import contextmanager
from html import escape
from urllib.parse import quote, urlencode
import sqlite3
import csv
import os
import json
import urllib.request
import urllib.error
import time
import shutil
import uuid
import mimetypes
from openpyxl import Workbook

application = FastAPI(title="Protean Booking System")

# Use absolute paths and a persistent data directory.
# On Render, attach a Persistent Disk at /var/data. The application will use it
# automatically when available. You can also override DATA_DIR, DB_FILE,
# UPLOAD_DIR and EXPORT_DIR through environment variables.
BASE_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = Path("/var/data") if Path("/var/data").exists() else BASE_DIR
DATA_DIR = Path(os.getenv("DATA_DIR", str(DEFAULT_DATA_DIR))).resolve()
DB_FILE = Path(os.getenv("DB_FILE", str(DATA_DIR / "protean_bookings.db"))).resolve()
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", str(DATA_DIR / "uploads"))).resolve()
EXPORT_DIR = Path(os.getenv("EXPORT_DIR", str(DATA_DIR / "exports"))).resolve()
BACKUP_DIR = Path(os.getenv("BACKUP_DIR", str(DATA_DIR / "backups"))).resolve()

# Configure these in the hosting environment instead of hard-coding production secrets.
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "Protean123%")
TEAMS_WEBHOOK_URL = os.getenv("TEAMS_WEBHOOK_URL", "")

# Microsoft Graph / SharePoint settings. Configure these as Render environment variables.
MS_TENANT_ID = os.getenv("MS_TENANT_ID", "")
MS_CLIENT_ID = os.getenv("MS_CLIENT_ID", "")
MS_CLIENT_SECRET = os.getenv("MS_CLIENT_SECRET", "")
SHAREPOINT_HOSTNAME = os.getenv("SHAREPOINT_HOSTNAME", "proteancoza.sharepoint.com")
SHAREPOINT_SITE_PATH = os.getenv("SHAREPOINT_SITE_PATH", "/sites/protean")
SHAREPOINT_FOLDER_PATH = os.getenv("SHAREPOINT_FOLDER_PATH", "General/Booked Claimants")
SHAREPOINT_DRIVE_NAME = os.getenv("SHAREPOINT_DRIVE_NAME", "Documents")
SHAREPOINT_REQUIRED = os.getenv("SHAREPOINT_REQUIRED", "false").lower() in {"1", "true", "yes"}
MAX_FILE_SIZE_MB = int(os.getenv("MAX_FILE_SIZE_MB", "25"))
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024
ALLOWED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".txt"
}

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)
BACKUP_DIR.mkdir(parents=True, exist_ok=True)
DB_FILE.parent.mkdir(parents=True, exist_ok=True)

LOGO_BASE64 = "iVBORw0KGgoAAAANSUhEUgAAAXIAAAB4CAMAAADYFonpAAAAAXNSR0IArs4c6QAAAARnQU1BAACxjwv8YQUAAAHgUExURQAAAL4YPb8QQL8aQL0XPb0XPboVOr8VQL0XPb0XPLwXPb8gQL4XPrwVPL4XPr0YPr0WPb0YPb0XPbwWPb0XPb0XPb0YPb0XPr0XPTxITTxHTDxITAAAADxITD1GTTxITEBAUDtHSz1GS0BIUL0XPbsYPL0XPTxHTDxHTDxFTDw8Sz1HTTxHTDtHTDxHTDlGTb8YQLsXOzxITDxGTDxHTTk5OTxITEBAQDpFSjlESTtJTTxHTDxHTDxHTDxHSr4XPUBKSjxHTDxGSz1ITT1HTDxGTT1HTTtHSz1HTTxHTDxITT1HRz1HTTtHS70XPL0WPTtGSj5GTTxHTThESzxITD1GTTxHSztGSTxHTTtGTL0ZQL0ZPz1ITT1ITTxITL0ZP70ZP7obQL0YQD1HTL4ZP7sXQD1HTbwaQL4ZP70ZP78gQDpGSztITT1ITL0ZP70ZPjxHTb4aP70QQj5ITj1HTTxHTTtHTDxHTBhgSBpgRiBgQBVgQBtgRBlfRBpeRBpeRBpeRBpdQxpeRBpcQxpeRBxgRBtfRBpeRBpfRCBgQDxHTBpeRBldRBtdRRlcQxtgRR5iRBpeRBtfRD1HTBtdRSBgRh9gRh5gRh9gRh9gRh9gRyBgSCBgRTpITL8YPj5GTMq4QF4AAACgdFJOUwC3ECj/3zAYj4dYCMdIz4Cfl/dQp69geOemv6cBgFBAEHBYIO9AcP7/URHfr2jvKCA4rse3CX8IMDE4z+fXSL8Y95hgl3ifgZ5h9hmPQddocSHmKe6giEl3eXCvzo6H5/8wYMbPON5Qv/cQaWeG74++xx9flraokCAoEBg4l3jvx1jnUP9Aj4CvCG+HcGhIMETPv9ZggLeHp/+/QGA5f1f2WXuGAAAACXBIWXMAABcRAAAXEQHKJvM/AAAYtElEQVR4Xu2di3/b1nXHj+krURatpy3ZJgmAImUbm0hRjhiD0YMiR4pzEs7pa2ua2oy6ZW2XNlu3ZFsTJ2mayEud9bE29bbs+a/uc84FcHEfAGibVLNP+ft84oi4AC7wxcG55557AQCcsc5l1CVTTVbn2cysvCQ75/+RuSAXTDUezTM2k4suyF1kC/yPRTa9ASahJcbYsrKA/15mTLH/qcaiFcZCtGjd2VXGVtHsLzHGLqlrT/XUyp7Hf5HsZUS+RgvPr64DXMHfCwC5q4yxeXI5mS+TsV/LF4rqsmeVZZDtlNTVotpQ1yeVnYq6YqhLm9cRdw7gHCLmZr7GrmZhfmlp7uIcwBxj7AasZAEyi0F7ClBSK0lTpM6KWmZU/DGTnJuu+wfqQk28qj9UF+tyLMsqgmvWVrUWi31LXTnQdt6qqyuT1tnmDsDqMsAFtrmyQ8vQk/hePUf+ZR3W2S1sRYWvr6g1pCncEsBRy4xyIlsY9JzrurupZo4XxnUbz6vLNd2+6bqxyFEFW92EKxY5arfmqesjcrYGGXTWC0tBWLLGGFvNBmtkryzDzia7BDeizevvGnmTWL6gLlbFkbvV1GuTitx1d/fUjVAh8v28r4PoRoctdQOy6MwOQ08S6Dx5mBvRtY7Q0i8ythIuaQcVRBTWdaiW5PP5yN5C5B11paiakS10oZGPgNJH7v5R2ooK8rKD6lpWtRMuc92ewWZD5BET2TuO2H5B3Qj53rrOGLscLqJ2ky1GVrrBGFuCTeU6aOoHtXTVElkh8i21ZGRxI0838wC5+8cpzBXkkYbE3g+Xug3dpZuQA0B3W2ykMM+tMnYdkWPDyXtCR4Q86P3k/JhlCRhj2NTG6wyR3/H30EshGSJ3X1SLZMUjl/Dta8xjkENFXKmqXIKxIUc+B3CZvAsnzm7x8nM57Jdy5FeVbRWdHfL2S/4eOi+rRbIE8pTwJgk51IWf6EsFCcihIlySwuN6gHwRYBPD8oyPfJ2Kb7BZChPJsYgY0aizQz4Iz2aQbOYR5J0/UQujSkQOe2GBW5ZL4pFDN9xmuy2XrLFZQs5gh5w1bz195LdW2Tr3NEswc1VKweg6M+ShkbtuJznmjiB3t5McfzJysMKSA8UzxyOHRrjRcXRxJpNZzHLKuQXGVs9DNoI8d5exCxiqMPYKHM0Cj9vjdGbIhZG77lfUQklR5O7WV9VioRTk3m5YpMTnCchFIL0bXZxdzGQAVhHqDgYmK4CRCQrTKujF12AGf96CDCzPRzfVdFbIS8LIXfcgsZsjIXcP42+JFOQR19KQCxKQg2hBJc+yupgBbsc7mEtczcIiR44WjX8eEfLVDGYUvxwRywY6lMNgJ4lmjsg7om+yH3t90pCD2InsWZKQi7tRKpxhF3PwCjGmZnIOo/C7i2zG7yexRUJ+mVKNlAOL1Rkh976G0WEtsN/DpAYUkW99/RtBhfFd/1Tk4cm5cuYkCbloAaRG9whbxhya+c4yEp6BzOoyZK7iQBAlF2cQ+eYOLKC7j26p6YyQo5G7LxB4UlL0R8iLjvBEfxpzgVKRC3xyFz4JuYhZJCKXGVudhexVxnYoY852YD0HcAlDdHLiZOXXIXP1SxKXE+tqkZNH7cdQRHHkUBYu/c/Ma6cij8GXiPw43EZqc7F7v5jD8DyLvpwxbCJz5EF47HIEd9kSt/iL0Q11nQ1y8igvAHihu/imuoqQjxxeFcy/ZWT+BMjlyDwJufDlUkaM/PU6Drv5yRXs77z27bCInYOjI4BZ/FNktYw6G+T3gnRWeEIJya0AOdwXzF80rZ6KfBiWyWyTkOeDsm2pyd3hzgNgBS2dZ8pzr598hyd2GcblFzI8oet3SGN1Jsht38ijHaL4Pk6I3E89ojqmuyIVeTWmLAF5O9ykIBeQw74EuRzkAuR/fnKCZs6RUxH3Mck9obNBjkbue+8RzFwg9/4iWNvtGC5RKvIwKD2QlycgFy2ukoimOGUZltZ5GI7I3zg5OfnLAHkWZi/T8D/FjUk6C+Rk5H6MUg+dRWxySyCH0neDtd3tr6vrpSIvhUmq6JBiIvJ6uImaSiSwV2GZLVFenF2AnZOTk5M3gMJCtgYLmzO8r5Tiys8EOdpqGKKEziI2uRVBDhURnutd/zTkoiVUUlSxyEth3/NA3VuGevuz5zBWQTNfh9cQ+cn3uJufzW6yTe5XkvueZ4KchibCQDwYqIjP4UaRQ1OE5/tq1z8FeTMsGSglccgF8Y4+gEeJwvkjxjazaNc78G1C/n1y85dp5J9il5So/EyQo11H4vDQzK0YM5eQQ0uELXmlG5qCPOS3rw6rmZF7G2GCYFsnznv7VzDunoN1tgi5vyLkJ2/CHFvG9CK7SuFjSrb8KZAbJPiYRGb9A/GbPDsqLrklI4cfCuZKk5uI3AvPbFubJ2FCHgHu7mtbYL8Hx9lWsOu5moNLr3BXfnJysgOzczw6vItNrBjzj9PkkaNVS3QxfiHFJLcU5PCWYC53/ZOQN8P4eku1fgPySrkngG9vKKv7wlZzmWyd5sH5yP+ayihwvIzIU4188sjJyKWTSDNzFTlcE1VJI9AK8t1hQLBUFhG5YbRZIDeo0VXdUCA082UaplgCHIh4nZD/DWQuZnmv85Vz0qyLOD0x8m11hpZlxVgFF4YNCtsw8jMntzTk3o+CDVz3B5ECBTmKpnxEhve39CkpScgPj5XYRtI8m7tC0coyzLLN8zxi+U5ukd2gQYrV7NryBT7vOVFPjPwJm0/qbirXZCMwc3NyS0MOpdAXSZfJgFxSJ3W2Fp86FD+ur2g5x+PzZezzLOZeP3n95Ps4ynzOXwopQ3Bck0aORt5RHEhKDldHDhXRJYp0/RXk5ZpVCPh18g3LifMQqi+PjNgl3rCQg0s72FCuUF90/bU33vzb72FUPkMT/TOZ0eY6Txg5GbkaFgszz5vM3IAcmqJLJEagk5rPRKnNZyWc9GKIxyUtsZW/u8iuEPJzuR3I5jAwnIFb7Orslc1z6upGTRg5GbnW50lObpmQQ0t0iQ6CbujYkIMd7mTL0NqGOo+J2rs769cperlKc7QwUL8MmQtvrmESIDfCoyuTRU4uRDNygPvB3kzJLSNyeFuEilt+N3R8yCO5gSR3Pk/d/os5v/sPcHmWOqXYZuKI6AI6nlRNFjmNAvXUCMeyesHeTGZuRk6Affld/zEiH82dX7pEz62so1NHK59la2jlc5CjRMsiwJG6iUETRS7ayVgZklsxyOEdwZyPQI8R+WjuPDeTzW5i1vDKPMxv3oU1tppbYRdgdpk8zQIspXeEJow8HOqMl+7oY5HD34vN/gHLx4k84s4P4mPzuSOMTRZhns1D9hWcPbSzcAvOb66g9a/Arc30jtBkkY9g5CYzj0XuRcLzfyyOGXnEnUfn1cvKzdyAdQpQVjFFy0eAsjMUwiwD3L2ibmDSJJGL8TddIoer9fpjkUNbhOc4Aj1e5F44hqSOaER0awbg1hoODOE46BEOBsEKprIurANc/90HiWiUsY8GhTlcLbkVjxzqIjx3XyyOF3lkREjNo0d0aRZ+zMeIFuDdBbb43oPsKvU7H7wP/BHEVE0QeXT8TReCJWnJrQTksBeZ2vhNa7zII3MwEtw5wAcf0lDo0o9/8mB94aOfXmc4XfHdjz9R14vTBJFHBplNis3hJiEXaUjX7WBgN07kAkaCOwf45PQh2vn8w9N/gk9OPz1/IwcP4GcfP1LXi9PkkCcbeQSeelkSkcNngjlqrMhHcucAH/wzfPohwLunP4Ofn/6CW/gH76trxWpyyDGkU2lKisvhJiOPTigaN/LR3Dnqk9Nf/urRB5/+6l9+/eDBex+MbuGoiSHXhyZUxeVwU5BHJhSNHfmo7hzgw/d/+dEj+PVv4N2PH/0UG9TRNTHk2vibJhG2y08EpSGPTCgaO/KoO4/N/fr6zS/ePz19+BE2pk+kSSFPN/LIzCg5uZWGHEqfBxuOH3nEnRvScZIePfzt6elDbEqfTJNCnm7ksTncVORQEd3acSOPuvOYMSWhh6cUvDyhJoS8jUYuPVNmkjmHm448MqFo7MijzyGm7fG93/78PXVZuiaE3Dw0oYouDEpadQTkYkLR+JFH3Plusjt//K+PH//bv6tLUzUZ5ObxN13GCYqjIA8nFE0A+cju/PEXX3zxxX+oS1M1GeSjGXnMBMWRkAcjFhNAPrI7Pyvko7wcZFQjj6TAI3OCRkMOd4j5JJCP6s7PCnmi/FOgR5y0uckmmWZujYjco8s1EeQjuvMvEfLg+bdRZEhujYicTygqQniLJXcXoyood6UmrxruVX3lQkRPifw42LVxJplQMzyGJPFTaIk/U0XrosSr5Zo/iv5KEE4oGmW9Sekpkf+/VvMbU+S/X5oiP3NNkZ+5/vPx48eP/0tdOtVUU0011VRTTTXVVFNNNdVUE1ZllKS9F5emn6DSqyzFPz2iCjO/hmT/QF/cjeaZtc1EenpoWXZc7hnLzYn59nGeHn7ZryaN3rUK/GUBW/0RzzC+wkjmO5BxTkWrQM/CbuetuP2Q7t/UXzJEupbP5+WnBHB36nuE/ZexKHM5carSQfQIJeTKIEzHzASLTJa8R+8nPmhYVuPQPYwb19jYcg+GewBeq7cdv5akuApROGAonZCO3LPoEud7Vj/vdvrx49Htl+jJFINw7qg8n44AafPC6L0JBuTau8xDIXIaS2+WLf69iIbBKmIIbOAIdTDMY++7fc0G6MMUrrsb7LO55bqFpIdMfcVUSELk8SeEquPIWeeY19MedA5iHmqh2bUxEwRikKtmzt849HTIUSUaoNzWh8qMBLwCnpjwJ17V3dVpOtvSRy7au657YHrxiyxjhb5Skdv41MmucBitl4yv4gtmCJjN3IC8gxammHnVRb/69MgBWmjoHY25kQDNIZG8f0F/ZBcfc5Re0O0dpDxMTTJW6CsNeR05DqNLWi91zB6bpu1vG2c0GpBv4bOOspk33X0002dBDk28kuq7+o0EaA6GXJd3oFbuYRsmN6x43MprvnWZKgyUgpzev6cM+1vhs+GSvK918J1BXzGZuQk5mots5lXXfmbk/F2iygsqTQTaeG3U93fZB8oNgnyVaVe0YVJ4gzJUGCoFuWkCl/f50NTMfHazj/NbjBN3TcjpdKJm3nT3abbIsyHnE+WUaNFAgOZsa+GNcmK6WfhM9pVlqgwVhkpGbp7AZQKO02BeoPkt0TcMBTIiV8+n6tpjQA413WEYCFDlh9Iig+hBb/W6lHFhipnrFQolI+cvMFeXGvXZzWqRpnGZnjQyIudmHrZEDprOGJBT2KNYoU6ghYu0xlIVBa1qw0BfbtDsUJZeoVAycnzwYduA0KDv0qXBGV+G5+nMyMnSwmPfR8uJQb4fvCxDOw0duYen25EWGQjQvSCFBSZh46nsyn+nccrV0isUQuTxJ4R+xfiOIU02GjmPWgxmbkbOP3XhG5FNlhmDPJSWdNCR89cUy2GcToCe/U/tSGJLqXsf81JJeoVCYrqqRsWfvN/XCRp0j79OniY2/rdaGIc8auZk5HHI4+9DA3J616LsDXQC1GnSLqAqXEmKykl41Cmzx/UKhRIdyx4ij3wcLnycXHsLq33Tt22cQK7PJY1BTrc3N3Nu5ONDLi/SCdDrfFKtHHtmGlzyXPqFkKRXKJSInKw8MreTvlyL/k1Dfi/w4BTjaA1uHHJylmTm3MjHghwnRSuLdALkrlJ9Oc2vVhfSd6PUyF+RXqFQInKjL8dPI6rIQyPnjzFrM3BjkeMrbNDMfSMfB3Jq3RQiOoHRIhZyPyo6Sl6mXC3TdoGSkWOcrUYsJuTSo8t65ykeeWDmvpGPAzm1TkrOXSfgoc9IaQP9C6N6fKogLjPvS69QKBk5PVihtIYG5M2bB0HQY1n4Jgo1uRWPnJt5YORjQO6hM1BTlAYC6MzlpBZK7X1i9KOG4Bisp90fhgpDJSMnz6x8XNWA/Lnoq0Lw4S3VzOORczMPjHwMyCn6U/voBgKUKlEbwZaaY9nrqB9xgTruTMtVKjJUGCoZOT8BGZaOvHlTSqygl1E+KZSAnL+QLHiTyjMjp66h9iC/iQDlv2rSotKBdrGwdnkZIlENX5OpwkApyOnB+0MpU6Ujl4ycm7nhrdAxyPlr+YM7/FmR2xQza+lsIwEchJNeTF/aNziMXWV/+Ald7ZJqMlboKwU5/5KHdFE15BXZyHmjK+dwk5CjmYdn+mzI2zSUaRiJMxIo4dodMfDY3nfz2sWC+q7r7ottmzhkoa+lylihrzTk4CDzRsQ3a8jvqy990nO4Scix6xI2Y0+F/IC+K18b5NE9H6iuARVDgL7Ovdsigp7VMY99wnHH3QoOsNVJjQ9JMRWSUpFDG5/U3A6iLv7xh29Fy1/SUuRo5lJyKxF5M3I7xyCXJJVTkCyUN7+tH4tMBNo4/Inj7T205bjOaH3X7fRbTdizC2J0OlnyUcnHLOVYUAb+ZQyVDgeO42wMEHhDShbf199TiZ9DlJJbiDyqIkQ/ElATDtW2LEuJgx0Rf3JJ5dQh9lVWI+hQWGr2Bk0Kat3tfCF2YzTuAeXKtvJD04UzKHJUpOjOm2qhMfde7/EJNq6bL9QUV/mqPrfFu21ZVjRO/EypRL1GU0011VRTTTXVVFNNNdVUvxcqlW2AkgNgl6Bkdalr2LbLOCZnWZjRqNi4zMYFfIvgC+LtMvYAm9hnK1mW1aUVaB2H/rVwJ9j9A2jWHNwJ/VMp+z31JX8/ntMNdm4DVFp8PfzNS6CO07TwEEnX+Wdv8TeWOgBOl770hVvYAJ5d9qjzXAGHOpsVPIsSJRHC7mVQd92mvETbLuOh1bt7AM0Krdc8pslhzTrvWNv8APxz3vN31cLDww1I/8M7ln5p6ba1UQxOhSgFZKqtMu4UoF+BSt8pVADqBcdpAvQdpw3QtZxCHaB/DOCnIHzkzZ4ztAG6G60CeE6/jFvwdXplG/NS+06/DtBwkM2GXdgDwPk5do9fkPC0S1XbwQrw7z6Ak/dovT5Am5eAhfkrPESSjxx/V4cAFl7bcqHN99DH2sslqDsNpwRW2e7hFcF8Cs3RDZMoft1Wzam2AerVllMBqA2cFkDXAbhTpEMuAgyHeHUbTh0qVhH2/HMuFHEVgDuYLcENSD7yboNKK285ewC3v0p7e3ujda0I8E7Z7heBMlkBcgvwAvR4RoofXx6gjShrbQU5Ui1QhfgXzXbiyJsWQNfGLSvBdeo64BX4SmJ6sn/aG/yIQ+TljQB58DXMAT7GYELeH7YJeYUfIt+BP8CNv6wK/iLkPcr7KMi9BsAe7ocyJaUGIQuQO3CnBN5wgDCQLyK/9jzAc0Xolm2+6HhYNCDfoNJKF38GyB248zw0rSL8bxHKCFhYOf4qQInsznFK0MYUfQMBD1Xk/L/uRgtXIeQNx2kAdFsAdQsg7w3LuGyPH1WfsxEpfx+5n6gNkTs9z0de5SV73WY3Bnnb8pEjWv+AWnyHhLxsWxx5aVjCihXke2ikjeAqIfwo8vq9IrRaNrrBEHkR8L9u5U4RF71a36gYkPPSylvoKkIrtwdFeLsFgJycgieQ1/AgClCx0GoQeQVv6jwe7BAtCaUi72GVEeQ1hwPZIv8VIq/SSnUxEhEgj/yPkDdrPnL/huiVoBeDHI49HzkdIv9vr4HMCXkXXRsix3tRR+7g8J+/JUCL51QD5H3EOfA8HHGOIseyit3CuXN3ipVjuh+4QuRUqiAfOAA/RDK4ynELPCRM5AclOgg8Cn4gVXyWB3+0h2bkDm4cdSzOBoDd5bvx94KOpRGaoy8fuZ+GFciB5mSiV+Yl+/n8YSkGebvGkdfRbdWhRCt1sankB4U3FSLHz2ryO42L113qA1QGwY3mmxQ2Ne/4IL17+TyOfgbInytyx1KB2wj183y+YULOS8mx4E1yHx2Ld60IDranAG0YtAEKHlVvwV4NwB56AnnPAatFP4b+AIaPvFyDFnptB4ZtCblXaJfQO8rIhy2+0rAMfsLcR16plnAKRQR5E8fbsFGgEjRFpxuDHI6xBa14/SZdZjz7NgywzfXtoOfhxngp2j0NOVTr0GtiFt6jQ7Dx31IPmhu+u2hhc9oSyLt4zoTcvleEV+sAtysm5FjKkVduF50u+fJ3Xgbv2sulO0XY6OHtX++hSVGsiBd6gLeBTQl9r9b3o6+2PyAVBIndXs2jEAmjMIr//CCxMhxgXMZXpzGJpmVhyEUrdfu+Ow8DtWGhCbCBwSQGiRg6+EEilWDYBDbGofzGjwSJeFA2jZ3Q2EqtilXy4+XnUfIPr47ltl9JpO7SRp94OYM+mkytgG7PLhzz8wJo4SVo8Z1RkOifcwng1SKgl6808eR4DBYEibyUB4lgP4f/NitQbwJU3hpocxZHUYD8WRUgf0L5yJ9NT1l3snzkE9AUeYwmh3yqZ9P/ARrtCZVP2oe9AAAAAElFTkSuQmCC"

CSS = """
<style>
body{font-family:Arial,sans-serif;background:#f4f6f8;margin:0;padding:20px;color:#111827}
.container{max-width:1700px;margin:auto;background:white;padding:22px;border-radius:12px;box-shadow:0 4px 18px rgba(0,0,0,.08)}
.login-box{max-width:480px;margin:60px auto;background:white;padding:28px;border-radius:12px;box-shadow:0 4px 18px rgba(0,0,0,.08)}
.brand-header{text-align:center;border-bottom:1px solid #e5e7eb;padding-bottom:18px;margin-bottom:20px}
.logo{max-width:430px;width:100%;height:auto;margin:auto;display:block}
.nav{display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap;margin-bottom:16px}
h1{text-align:center;margin-bottom:6px}.subtitle{text-align:center;color:#555;margin-bottom:20px}
.top-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:15px;margin-bottom:20px}
.bottom-grid{display:grid;grid-template-columns:1fr 1fr;gap:15px;margin-top:20px}
.table-wrap{overflow-x:auto}table{width:100%;border-collapse:collapse;font-size:12px;background:white}
th,td{border:1px solid #d1d5db;padding:6px;vertical-align:top;word-break:break-word}
th{background:#1f2937;color:white;position:sticky;top:0;z-index:1}
input,select,textarea{width:100%;box-sizing:border-box;padding:6px;border:1px solid #c7c7c7;border-radius:4px;font-size:12px}
textarea{min-height:70px}.badge{background:#e5e7eb;border-radius:999px;padding:4px 8px;font-size:12px}
button,.btn{background:#1f2937;color:white;padding:10px 16px;border:none;border-radius:6px;text-decoration:none;cursor:pointer;font-size:14px;display:inline-block}
.btn-green{background:#15803d}.btn-orange{background:#c2410c}.btn-red{background:#b91c1c}
.notice{background:#ecfeff;border-left:4px solid #0891b2;padding:10px;margin:12px 0}
.terms{background:#f9fafb;border:1px solid #d1d5db;border-radius:8px;padding:14px;margin-top:18px;line-height:1.5}
.error{background:#fee2e2;border-left:4px solid #b91c1c;padding:10px;margin:12px 0}
.success{background:#dcfce7;border-left:4px solid #15803d;padding:10px;margin:12px 0}
.actions{display:flex;gap:10px;flex-wrap:wrap;margin-top:15px}
.check-row{display:flex;gap:10px;align-items:flex-start;background:#fff7ed;border-left:4px solid #f97316;padding:12px;margin-top:15px}
.check-row input{width:auto;margin-top:3px}.file-list{margin:0;padding-left:18px}.file-list li{margin:3px 0}
.small{font-size:12px;color:#555}.nowrap{white-space:nowrap}
.upload-box{border:2px dashed #9ca3af;border-radius:10px;padding:16px;background:#fff}
.upload-box input[type=file]{padding:10px;background:#f9fafb}
.selected-files{margin-top:12px;border:1px solid #d1d5db;border-radius:8px;background:white;padding:10px}
.selected-files-title{font-weight:700;margin-bottom:8px}
.selected-files ul{list-style:none;margin:0;padding:0}
.selected-files li{display:flex;justify-content:space-between;align-items:center;gap:12px;padding:8px;border-bottom:1px solid #e5e7eb}
.selected-files li:last-child{border-bottom:none}
.file-meta{min-width:0;word-break:break-word}
.remove-file{background:#b91c1c;padding:5px 9px;font-size:12px;flex:0 0 auto}
.empty-files{color:#6b7280;font-size:12px}
@media(max-width:900px){.top-grid,.bottom-grid{grid-template-columns:1fr}}
</style>
"""


FILE_UPLOAD_SCRIPT = r"""
<script>
(function () {
    const input = document.getElementById("document_upload");
    const list = document.getElementById("selected_document_list");
    const count = document.getElementById("selected_document_count");
    if (!input || !list || !count || typeof DataTransfer === "undefined") return;

    const selectedFiles = new DataTransfer();

    function formatBytes(bytes) {
        if (bytes < 1024) return bytes + " B";
        if (bytes < 1024 * 1024) return (bytes / 1024).toFixed(1) + " KB";
        return (bytes / (1024 * 1024)).toFixed(1) + " MB";
    }

    function fileKey(file) {
        return [file.name, file.size, file.lastModified].join("|");
    }

    function refreshInput() {
        input.files = selectedFiles.files;
        count.textContent = selectedFiles.files.length +
            (selectedFiles.files.length === 1 ? " document selected" : " documents selected");

        list.innerHTML = "";
        if (selectedFiles.files.length === 0) {
            list.innerHTML = '<li class="empty-files">No documents selected yet.</li>';
            return;
        }

        Array.from(selectedFiles.files).forEach(function (file, index) {
            const item = document.createElement("li");
            const meta = document.createElement("div");
            meta.className = "file-meta";
            meta.textContent = file.name + " (" + formatBytes(file.size) + ")";

            const remove = document.createElement("button");
            remove.type = "button";
            remove.className = "remove-file";
            remove.textContent = "Remove";
            remove.addEventListener("click", function () {
                const replacement = new DataTransfer();
                Array.from(selectedFiles.files).forEach(function (existing, existingIndex) {
                    if (existingIndex !== index) replacement.items.add(existing);
                });
                selectedFiles.items.clear();
                Array.from(replacement.files).forEach(function (remaining) {
                    selectedFiles.items.add(remaining);
                });
                refreshInput();
            });

            item.appendChild(meta);
            item.appendChild(remove);
            list.appendChild(item);
        });
    }

    input.addEventListener("change", function () {
        const existingKeys = new Set(Array.from(selectedFiles.files).map(fileKey));
        Array.from(input.files).forEach(function (file) {
            const key = fileKey(file);
            if (!existingKeys.has(key)) {
                selectedFiles.items.add(file);
                existingKeys.add(key);
            }
        });
        refreshInput();
    });

    refreshInput();
})();
</script>
"""

@contextmanager
def db_connection():
    conn = sqlite3.connect(DB_FILE, timeout=30)
    try:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("PRAGMA journal_mode = WAL")
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def backup_database():
    """Create an atomic SQLite backup and retain the latest 30 copies."""
    if not DB_FILE.exists():
        return
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = BACKUP_DIR / f"protean_bookings_{timestamp}.db"
    source = sqlite3.connect(DB_FILE, timeout=30)
    target = sqlite3.connect(backup_path)
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()
    backups = sorted(BACKUP_DIR.glob("protean_bookings_*.db"), key=lambda x: x.stat().st_mtime, reverse=True)
    for old_backup in backups[30:]:
        old_backup.unlink(missing_ok=True)


def ensure_column(conn, table_name: str, column_name: str, definition: str):
    columns = {row[1] for row in conn.execute(f"PRAGMA table_info({table_name})").fetchall()}
    if column_name not in columns:
        conn.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}")


def init_db():
    with db_connection() as conn:
        conn.execute("""
        CREATE TABLE IF NOT EXISTS bookings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT,
            law_firm TEXT,
            law_firm_contact_person TEXT,
            law_firm_phone TEXT,
            law_firm_email TEXT,
            assessment_place TEXT,
            assessment_date TEXT,
            claimant_name TEXT,
            date_of_birth TEXT,
            gender TEXT,
            preferred_language TEXT,
            contact_number TEXT,
            occupation_status TEXT,
            claim_type TEXT,
            mandatory_documents_submitted TEXT,
            injuries_sustained TEXT,
            prescribing_date TEXT,
            protean_experts TEXT,
            additional_information TEXT,
            permission_to_contact TEXT,
            expert_affidavits TEXT,
            terms_accepted TEXT,
            external_experts TEXT,
            assigned_caller TEXT,
            call_attempted TEXT,
            date_of_call TEXT,
            contact_outcome TEXT,
            documents_requested TEXT,
            documents_expected_on_day TEXT,
            claimants_readiness_notes TEXT,
            booking_status TEXT,
            mandatory_los_documents TEXT
        )
        """)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS booking_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            submission_reference TEXT NOT NULL,
            booking_id INTEGER,
            original_filename TEXT NOT NULL,
            stored_filename TEXT NOT NULL,
            content_type TEXT,
            file_size INTEGER NOT NULL DEFAULT 0,
            uploaded_at TEXT NOT NULL,
            FOREIGN KEY (booking_id) REFERENCES bookings(id) ON DELETE CASCADE
        )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_booking_id ON booking_documents(booking_id)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_submission_reference ON booking_documents(submission_reference)")
        ensure_column(conn, "bookings", "is_archived", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "bookings", "archived_at", "TEXT")
        ensure_column(conn, "booking_documents", "is_archived", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(conn, "booking_documents", "sharepoint_url", "TEXT")
        ensure_column(conn, "booking_documents", "sharepoint_item_id", "TEXT")
        ensure_column(conn, "booking_documents", "sharepoint_status", "TEXT NOT NULL DEFAULT 'Pending'")
        ensure_column(conn, "booking_documents", "sharepoint_error", "TEXT")


init_db()
backup_database()


def option_tags(items):
    return "".join(f"<option>{escape(x)}</option>" for x in items)


def is_admin(request: Request):
    return request.query_params.get("admin_key") == ADMIN_PASSWORD


def safe_text(value):
    return escape("" if value is None else str(value))


def fetch_bookings():
    with db_connection() as conn:
        cur = conn.execute("SELECT * FROM bookings WHERE COALESCE(is_archived, 0)=0 ORDER BY id DESC")
        all_headers = [d[0] for d in cur.description]
        all_rows = cur.fetchall()
    hidden = {"is_archived", "archived_at"}
    keep_indexes = [i for i, name in enumerate(all_headers) if name not in hidden]
    headers = [all_headers[i] for i in keep_indexes]
    rows = [tuple(row[i] for i in keep_indexes) for row in all_rows]
    return headers, rows


def fetch_documents_for_bookings(booking_ids: List[int]):
    if not booking_ids:
        return {}
    placeholders = ",".join("?" for _ in booking_ids)
    with db_connection() as conn:
        rows = conn.execute(
            f"""SELECT id, booking_id, original_filename, content_type, file_size, uploaded_at,
                       sharepoint_url, sharepoint_status, sharepoint_error
                FROM booking_documents
                WHERE booking_id IN ({placeholders})
                ORDER BY id""",
            booking_ids,
        ).fetchall()
    result = {}
    for (document_id, booking_id, original_name, content_type, file_size, uploaded_at,
         sharepoint_url, sharepoint_status, sharepoint_error) in rows:
        result.setdefault(booking_id, []).append({
            "id": document_id,
            "name": original_name,
            "content_type": content_type,
            "size": file_size,
            "uploaded_at": uploaded_at,
            "sharepoint_url": sharepoint_url or "",
            "sharepoint_status": sharepoint_status or "Pending",
            "sharepoint_error": sharepoint_error or "",
        })
    return result


def fetch_all_documents():
    with db_connection() as conn:
        rows = conn.execute(
            """SELECT d.id, d.booking_id, d.submission_reference,
                      d.original_filename, d.content_type, d.file_size,
                      d.uploaded_at, b.claimant_name, b.law_firm
               FROM booking_documents d
               LEFT JOIN bookings b ON b.id = d.booking_id
               ORDER BY d.id DESC"""
        ).fetchall()
    return rows


def format_size(size):
    size = int(size or 0)
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / (1024 * 1024):.1f} MB"


def clean_filename(filename: str) -> str:
    # Path.name blocks path traversal; remaining characters are reduced to a safe set.
    name = Path(filename or "document").name.strip()
    stem = "".join(ch if ch.isalnum() or ch in "._- " else "_" for ch in name)
    stem = stem.strip(" .") or "document"
    return stem[:180]


async def save_uploaded_file(upload: UploadFile, submission_reference: str):
    original_name = clean_filename(upload.filename or "document")
    extension = Path(original_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"File type '{extension or 'unknown'}' is not allowed: {original_name}")

    stored_name = f"{submission_reference}_{uuid.uuid4().hex}{extension}"
    destination = UPLOAD_DIR / stored_name
    total = 0
    try:
        with destination.open("wb") as output:
            while True:
                chunk = await upload.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_FILE_SIZE_BYTES:
                    raise ValueError(f"{original_name} exceeds the {MAX_FILE_SIZE_MB} MB limit.")
                output.write(chunk)
        if total == 0:
            destination.unlink(missing_ok=True)
            raise ValueError(f"{original_name} is empty.")
        return {
            "original_filename": original_name,
            "stored_filename": stored_name,
            "content_type": upload.content_type or mimetypes.guess_type(original_name)[0] or "application/octet-stream",
            "file_size": total,
            "path": destination,
        }
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        await upload.close()


def sharepoint_is_configured():
    return all([MS_TENANT_ID, MS_CLIENT_ID, MS_CLIENT_SECRET])


def _http_error_message(exc: Exception) -> str:
    if isinstance(exc, urllib.error.HTTPError):
        try:
            body = exc.read().decode("utf-8", errors="replace")
        except Exception:
            body = ""
        return f"HTTP {exc.code} {exc.reason}: {body[:2000]}"
    return str(exc)


def graph_request(url: str, token: str, method: str = "GET", data=None, headers=None, timeout: int = 60):
    request_headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/json",
    }
    if headers:
        request_headers.update(headers)
    request = urllib.request.Request(url, data=data, headers=request_headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read()
            if not raw:
                return {}
            return json.loads(raw.decode("utf-8"))
    except Exception as exc:
        raise RuntimeError(_http_error_message(exc)) from exc


def get_graph_access_token():
    if not sharepoint_is_configured():
        missing = [name for name, value in {
            "MS_TENANT_ID": MS_TENANT_ID,
            "MS_CLIENT_ID": MS_CLIENT_ID,
            "MS_CLIENT_SECRET": MS_CLIENT_SECRET,
        }.items() if not value]
        raise RuntimeError("Missing Render environment variables: " + ", ".join(missing))

    token_url = f"https://login.microsoftonline.com/{MS_TENANT_ID}/oauth2/v2.0/token"
    body = urlencode({
        "client_id": MS_CLIENT_ID,
        "client_secret": MS_CLIENT_SECRET,
        "scope": "https://graph.microsoft.com/.default",
        "grant_type": "client_credentials",
    }).encode("utf-8")
    request = urllib.request.Request(
        token_url,
        data=body,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        raise RuntimeError("Microsoft sign-in failed: " + _http_error_message(exc)) from exc

    token = payload.get("access_token")
    if not token:
        raise RuntimeError("Microsoft Graph did not return an access token.")
    return token


def resolve_sharepoint_site_and_drive(token: str):
    site_lookup_url = (
        f"https://graph.microsoft.com/v1.0/sites/"
        f"{SHAREPOINT_HOSTNAME}:{SHAREPOINT_SITE_PATH}"
    )
    site = graph_request(site_lookup_url, token)
    site_id = site.get("id")
    if not site_id:
        raise RuntimeError("The SharePoint site could not be resolved.")

    drives = graph_request(
        f"https://graph.microsoft.com/v1.0/sites/{quote(site_id, safe='')}/drives",
        token,
    ).get("value", [])

    requested = SHAREPOINT_DRIVE_NAME.strip().lower()
    drive = next((d for d in drives if str(d.get("name", "")).strip().lower() == requested), None)
    if drive is None:
        drive = next((d for d in drives if str(d.get("name", "")).strip().lower() in {"documents", "shared documents"}), None)
    if drive is None and drives:
        drive = drives[0]
    if drive is None:
        raise RuntimeError("No SharePoint document library was found for this site.")

    return site_id, drive["id"], drive.get("name", "")


def ensure_sharepoint_folder(token: str, drive_id: str, folder_path: str):
    parent_id = "root"
    current_parts = []
    for part in [p.strip() for p in folder_path.strip("/").split("/") if p.strip()]:
        current_parts.append(part)
        encoded = "/".join(quote(p, safe="") for p in current_parts)
        lookup_url = f"https://graph.microsoft.com/v1.0/drives/{quote(drive_id, safe='')}/root:/{encoded}"
        try:
            item = graph_request(lookup_url, token)
            parent_id = item.get("id", parent_id)
            continue
        except RuntimeError as exc:
            if "HTTP 404" not in str(exc):
                raise

        create_url = f"https://graph.microsoft.com/v1.0/drives/{quote(drive_id, safe='')}/items/{quote(parent_id, safe='')}/children"
        payload = json.dumps({
            "name": part,
            "folder": {},
            "@microsoft.graph.conflictBehavior": "fail",
        }).encode("utf-8")
        try:
            item = graph_request(
                create_url,
                token,
                method="POST",
                data=payload,
                headers={"Content-Type": "application/json"},
            )
        except RuntimeError as exc:
            # Another request may have created the folder between lookup and create.
            if "HTTP 409" in str(exc):
                item = graph_request(lookup_url, token)
            else:
                raise
        parent_id = item.get("id", parent_id)


def upload_file_to_sharepoint(file_info, submission_reference: str, token: str, drive_id: str):
    folder = SHAREPOINT_FOLDER_PATH.strip("/")
    remote_name = clean_filename(f"{submission_reference}_{file_info['original_filename']}")
    encoded_path = "/".join(quote(part, safe="") for part in (folder + "/" + remote_name).split("/"))
    upload_url = (
        f"https://graph.microsoft.com/v1.0/drives/{quote(drive_id, safe='')}/"
        f"root:/{encoded_path}:/content?@microsoft.graph.conflictBehavior=rename"
    )
    data = file_info["path"].read_bytes()

    last_error = None
    for attempt in range(1, 4):
        try:
            result = graph_request(
                upload_url,
                token,
                method="PUT",
                data=data,
                headers={"Content-Type": file_info["content_type"] or "application/octet-stream"},
                timeout=180,
            )
            return {
                "sharepoint_item_id": result.get("id", ""),
                "sharepoint_url": result.get("webUrl", ""),
                "sharepoint_status": "Uploaded",
                "sharepoint_error": "",
            }
        except Exception as exc:
            last_error = exc
            if attempt < 3:
                time.sleep(2 ** attempt)
    raise RuntimeError(f"Upload failed after 3 attempts: {last_error}")


def sync_files_to_sharepoint(saved_files, submission_reference: str):
    if not saved_files:
        return False
    if not sharepoint_is_configured():
        # Do not fail or reject the client's booking. Files remain safely stored
        # in persistent local storage and are marked for later SharePoint sync.
        for file_info in saved_files:
            file_info.setdefault("sharepoint_item_id", "")
            file_info.setdefault("sharepoint_url", "")
            file_info["sharepoint_status"] = "Pending configuration"
            file_info["sharepoint_error"] = ""
        print(
            f"SHAREPOINT_SYNC_PENDING [{submission_reference}]: "
            "Microsoft Graph credentials are not configured; files retained locally."
        )
        return False

    token = get_graph_access_token()
    site_id, drive_id, drive_name = resolve_sharepoint_site_and_drive(token)
    ensure_sharepoint_folder(token, drive_id, SHAREPOINT_FOLDER_PATH)

    try:
        for file_info in saved_files:
            result = upload_file_to_sharepoint(file_info, submission_reference, token, drive_id)
            result["sharepoint_error"] = ""
            file_info.update(result)
        print(
            f"SharePoint sync successful: site={site_id}, drive={drive_name} ({drive_id}), "
            f"folder={SHAREPOINT_FOLDER_PATH}, files={len(saved_files)}"
        )
        return True
    except Exception as exc:
        error_text = str(exc)[:2000]
        for file_info in saved_files:
            file_info.setdefault("sharepoint_item_id", "")
            file_info.setdefault("sharepoint_url", "")
            file_info["sharepoint_status"] = "Failed"
            file_info["sharepoint_error"] = error_text
        print(f"SHAREPOINT_SYNC_ERROR: {error_text}")
        raise RuntimeError(f"SharePoint upload failed: {error_text}") from exc


def notify_teams(law_firm, assessment_date, saved, document_count):
    if not TEAMS_WEBHOOK_URL:
        return
    try:
        payload = {
            "text": (
                "New Protean booking submitted. "
                f"Law Firm: {law_firm} | Assessment Date: {assessment_date} | "
                f"Claimants Saved: {saved} | Documents Uploaded: {document_count}"
            )
        }
        data = json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(
            TEAMS_WEBHOOK_URL,
            data=data,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        urllib.request.urlopen(req, timeout=5)
    except Exception:
        # Booking submission must not fail because Teams notification failed.
        pass


GENDERS = ["Male", "Female", "Other"]
LANGUAGES = ["English", "IsiZulu", "IsiXhosa", "Sesotho", "Setswana", "Sepedi", "Tshivenda", "XiTsonga", "Afrikaans", "Shona", "Ndebele", "Other"]
OCCUPATIONS = ["Employed", "Self-employed", "Student/scholar", "Unemployed", "Minor child", "Pensioner", "Other"]
YES_NO = ["Yes", "No"]


@application.get("/")
async def root():
    return RedirectResponse("/client", status_code=303)


@application.get("/health")
async def health():
    return {
        "status": "ok",
        "data_directory": str(DATA_DIR),
        "database": str(DB_FILE),
        "upload_directory": str(UPLOAD_DIR),
        "upload_directory_exists": UPLOAD_DIR.exists(),
        "persistent_disk_active": str(DATA_DIR).startswith("/var/data"),
        "backup_directory": str(BACKUP_DIR),
        "database_backups": len(list(BACKUP_DIR.glob("protean_bookings_*.db"))),
    }


@application.get("/client", response_class=HTMLResponse)
async def client_interface(error: str = ""):
    rows = ""
    for i in range(1, 11):
        rows += f"""
        <tr>
            <td>{i}</td>
            <td><input name="claimant_name" placeholder="Claimant Name"></td>
            <td><input type="date" name="date_of_birth"></td>
            <td><select name="gender"><option></option>{option_tags(GENDERS)}</select></td>
            <td><select name="preferred_language"><option></option>{option_tags(LANGUAGES)}</select></td>
            <td><input name="contact_number" placeholder="Phone / WhatsApp"></td>
            <td><select name="occupation_status"><option></option>{option_tags(OCCUPATIONS)}</select></td>
            <td><input name="claim_type" placeholder="LOS, LOE, Medical negligence, etc."></td>
            <td><input name="mandatory_documents_submitted" placeholder="LOI, ID, Hospital records, RAF 1, RAF 4, etc."></td>
            <td><input name="injuries_sustained" placeholder="Head, spinal, fracture, etc."></td>
            <td><input type="date" name="prescribing_date"></td>
            <td><input name="protean_experts" placeholder="Expert(s) scheduled"></td>
        </tr>
        """
    error_html = f"<div class='error'>{safe_text(error)}</div>" if error else ""
    return f"""
    <html><head><title>Protean Booking System</title>{CSS}</head>
    <body><div class="container">
    <div class="brand-header">
        <img class="logo" src="data:image/png;base64,{LOGO_BASE64}" alt="Protean Medico Legal">
        <h1>Protean Booking System</h1>
        <p class="subtitle">Medico-Legal Assessment Booking Platform</p>
    </div>
    {error_html}
    <form action="/submit-bulk" method="post" enctype="multipart/form-data">
    <div class="top-grid">
        <div><label><b>Law Firm</b></label><input name="law_firm" placeholder="Law Firm" required></div>
        <div><label><b>Law Firm Contact Person</b></label><input name="law_firm_contact_person" placeholder="Contact Person" required></div>
        <div><label><b>Law Firm Contact Number</b></label><input name="law_firm_phone" placeholder="Contact Number" required></div>
    </div>
    <div class="top-grid">
        <div><label><b>Law Firm Email</b></label><input type="email" name="law_firm_email" placeholder="Email Address" required></div>
        <div><label><b>Assessment Place</b></label><input name="assessment_place" placeholder="Assessment Place" required></div>
        <div><label><b>Assessment Date</b></label><input type="date" name="assessment_date" required></div>
    </div>

    <div class="table-wrap"><table>
    <tr>
    <th>#</th><th>Claimant Name</th><th>DOB</th><th>Gender</th><th>Preferred Language</th><th>Contact Number</th><th>Occupation Status</th><th>Type of Claim</th><th>Documents Submitted</th><th>Injuries Sustained</th><th>Prescribing Date</th><th>Protean Expert(s) Scheduled</th>
    </tr>{rows}</table></div>

    <div class="bottom-grid">
        <div><label><b>Additional Information</b></label><textarea name="additional_information" placeholder="Please provide any other relevant details for any claimant."></textarea></div>
        <div>
            <label><b>Do you grant us permission to contact the claimant if additional information is required?</b></label>
            <select name="permission_to_contact"><option></option>{option_tags(YES_NO)}</select><br><br>
            <label><b>Would you like expert affidavits submitted with the reports?</b></label>
            <select name="expert_affidavits"><option></option>{option_tags(YES_NO)}</select>
        </div>
    </div>

    <div class="terms">
        <h3>Upload Supporting Documents</h3>
        <p>Upload the Letter of Instruction, ID, hospital records, RAF 1, RAF 4 and other supporting documents. You may select several files at once or select additional files again.</p>
        <div class="upload-box">
            <input id="document_upload" type="file" name="documents" multiple accept=".pdf,.doc,.docx,.xls,.xlsx,.csv,.png,.jpg,.jpeg,.tif,.tiff,.txt">
            <p class="small">Maximum {MAX_FILE_SIZE_MB} MB per file. Uploaded documents are linked to each claimant in this submission and become available in the admin backend.</p>
            <div class="selected-files" aria-live="polite">
                <div id="selected_document_count" class="selected-files-title">0 documents selected</div>
                <ul id="selected_document_list"><li class="empty-files">No documents selected yet.</li></ul>
            </div>
        </div>
    </div>

    <div class="terms">
        <h3>Terms and Conditions</h3>
        <p>By booking an assessment appointment, I acknowledge and agree to the following:</p>
        <ul>
            <li>All mandatory documents, including the Letter of Instruction, ID document, hospital records, and RAF 1 and RAF 4 forms, must be submitted using the same email address through which the booking link was sent.</li>
            <li>If the required documents cannot be submitted electronically before the assessment, the original documents or copies thereof must be brought on the day of the assessment.</li>
            <li>It remains the responsibility of the attorney to ensure that all supporting documentation required to finalize the report is submitted timeously and accurately.</li>
        </ul>
    </div>
    <label class="check-row">
        <input type="checkbox" name="terms_accepted" value="Accepted" required>
        <span>I have read and accepted the Terms and Conditions.</span>
    </label>
    <div class="actions"><button type="submit">Submit Booking and Upload Documents</button></div>
    </form>
    {FILE_UPLOAD_SCRIPT}
    </div></body></html>
    """


@application.post("/submit-bulk", response_class=HTMLResponse)
async def submit_bulk(
    law_firm: str = Form(""), law_firm_contact_person: str = Form(""), law_firm_phone: str = Form(""),
    law_firm_email: str = Form(""), assessment_place: str = Form(""), assessment_date: str = Form(""),
    additional_information: str = Form(""), permission_to_contact: str = Form(""), expert_affidavits: str = Form(""),
    terms_accepted: str = Form("Not accepted"),
    claimant_name: List[str] = Form(default=[]), date_of_birth: List[str] = Form(default=[]), gender: List[str] = Form(default=[]),
    preferred_language: List[str] = Form(default=[]), contact_number: List[str] = Form(default=[]), occupation_status: List[str] = Form(default=[]),
    claim_type: List[str] = Form(default=[]), mandatory_documents_submitted: List[str] = Form(default=[]), injuries_sustained: List[str] = Form(default=[]),
    prescribing_date: List[str] = Form(default=[]), protean_experts: List[str] = Form(default=[]),
    documents: List[UploadFile] = File(default=[]),
):
    saved_files = []
    submission_reference = datetime.now().strftime("%Y%m%d%H%M%S") + "_" + uuid.uuid4().hex[:8]
    try:
        names = [(name or "").strip() for name in claimant_name]
        if not any(names):
            return RedirectResponse("/client?error=" + quote("Please enter at least one claimant name."), status_code=303)
        if terms_accepted != "Accepted":
            return RedirectResponse("/client?error=" + quote("You must accept the Terms and Conditions."), status_code=303)

        # Save files first. If any file is invalid, all files from this submission are removed.
        for upload in documents:
            if upload and upload.filename:
                saved_files.append(await save_uploaded_file(upload, submission_reference))

        # Save the booking and local document records first. SharePoint sync happens
        # afterwards so a Microsoft outage never destroys the client's submission.
        for file_info in saved_files:
            file_info["sharepoint_status"] = "Pending"
            file_info["sharepoint_error"] = ""
            file_info["sharepoint_url"] = ""
            file_info["sharepoint_item_id"] = ""

        booking_ids = []
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with db_connection() as conn:
            for i, cn in enumerate(names[:10]):
                if not cn:
                    continue
                row = (
                    now, law_firm.strip(), law_firm_contact_person.strip(), law_firm_phone.strip(), law_firm_email.strip(),
                    assessment_place.strip(), assessment_date.strip(), cn,
                    date_of_birth[i] if i < len(date_of_birth) else "",
                    gender[i] if i < len(gender) else "",
                    preferred_language[i] if i < len(preferred_language) else "",
                    contact_number[i] if i < len(contact_number) else "",
                    occupation_status[i] if i < len(occupation_status) else "",
                    claim_type[i] if i < len(claim_type) else "",
                    mandatory_documents_submitted[i] if i < len(mandatory_documents_submitted) else "",
                    injuries_sustained[i] if i < len(injuries_sustained) else "",
                    prescribing_date[i] if i < len(prescribing_date) else "",
                    protean_experts[i] if i < len(protean_experts) else "",
                    additional_information.strip(), permission_to_contact, expert_affidavits,
                    terms_accepted, "", "", "", "", "", "", "", "", "New", ""
                )
                cur = conn.execute("""
                INSERT INTO bookings (
                    created_at, law_firm, law_firm_contact_person, law_firm_phone, law_firm_email,
                    assessment_place, assessment_date, claimant_name, date_of_birth, gender,
                    preferred_language, contact_number, occupation_status, claim_type,
                    mandatory_documents_submitted, injuries_sustained, prescribing_date, protean_experts,
                    additional_information, permission_to_contact, expert_affidavits, terms_accepted,
                    external_experts, assigned_caller, call_attempted, date_of_call, contact_outcome,
                    documents_requested, documents_expected_on_day, claimants_readiness_notes,
                    booking_status, mandatory_los_documents
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, row)
                booking_ids.append(cur.lastrowid)

            # Link every uploaded document to every claimant in the same batch.
            for booking_id in booking_ids:
                for file_info in saved_files:
                    conn.execute("""
                    INSERT INTO booking_documents (
                        submission_reference, booking_id, original_filename, stored_filename,
                        content_type, file_size, uploaded_at, sharepoint_url,
                        sharepoint_item_id, sharepoint_status, sharepoint_error
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        submission_reference, booking_id, file_info["original_filename"],
                        file_info["stored_filename"], file_info["content_type"],
                        file_info["file_size"], now, file_info.get("sharepoint_url", ""),
                        file_info.get("sharepoint_item_id", ""),
                        file_info.get("sharepoint_status", "Pending"),
                        file_info.get("sharepoint_error", "")
                    ))

        sharepoint_ok = True
        sharepoint_message = ""
        if saved_files:
            try:
                synced = sync_files_to_sharepoint(saved_files, submission_reference)
                with db_connection() as conn:
                    for file_info in saved_files:
                        conn.execute(
                            """UPDATE booking_documents
                               SET sharepoint_url=?, sharepoint_item_id=?, sharepoint_status=?, sharepoint_error=?
                               WHERE submission_reference=? AND stored_filename=?""",
                            (file_info.get("sharepoint_url", ""), file_info.get("sharepoint_item_id", ""),
                             file_info.get("sharepoint_status", "Pending configuration"),
                             file_info.get("sharepoint_error", ""),
                             submission_reference, file_info["stored_filename"]),
                        )
                if synced:
                    sharepoint_message = f"{len(saved_files)} document(s) uploaded to SharePoint."
                else:
                    sharepoint_ok = False
                    sharepoint_message = (
                        f"{len(saved_files)} document(s) were received and stored safely. "
                        "SharePoint synchronisation is pending administrator configuration."
                    )
            except Exception as sync_exc:
                sharepoint_ok = False
                error_text = str(sync_exc)[:2000]
                with db_connection() as conn:
                    conn.execute(
                        """UPDATE booking_documents
                           SET sharepoint_status='Pending retry', sharepoint_error=?
                           WHERE submission_reference=?""",
                        (error_text, submission_reference),
                    )
                print(f"SHAREPOINT_SYNC_ERROR [{submission_reference}]: {error_text}")
                sharepoint_message = (
                    f"{len(saved_files)} document(s) were safely retained and queued for SharePoint retry. "
                    "The booking was not lost."
                )
                if SHAREPOINT_REQUIRED:
                    print("SHAREPOINT_REQUIRED is enabled, but the booking remains preserved locally.")

        backup_database()
        notify_teams(law_firm, assessment_date, len(booking_ids), len(saved_files))
        status_class = "success" if sharepoint_ok else "notice"
        return f"""
        <html><head>{CSS}</head><body><div class='container'>
        <div class='brand-header'>
            <img class='logo' src='data:image/png;base64,{LOGO_BASE64}' alt='Protean Medico Legal'>
            <h1>{len(booking_ids)} booking(s) submitted successfully.</h1>
            <div class='{status_class}'>{safe_text(sharepoint_message)}</div>
        </div>
        <div class='actions'><a class='btn' href='/client'>Submit Another Booking</a></div>
        </div></body></html>
        """
    except ValueError as exc:
        for file_info in saved_files:
            file_info["path"].unlink(missing_ok=True)
        return RedirectResponse("/client?error=" + quote(str(exc)), status_code=303)
    except Exception as exc:
        # Keep local copies for recovery if SharePoint or another external service fails.
        # Avoid exposing server internals to clients; log details in hosting logs.
        print(f"Submission error: {type(exc).__name__}: {exc}")
        return RedirectResponse(
            "/client?error=" + quote("The booking or SharePoint upload could not be completed. Your locally received files were retained for recovery. Please contact the administrator."),
            status_code=303,
        )


@application.get("/admin-login", response_class=HTMLResponse)
async def admin_login(error: str = ""):
    message = "<div class='error'>Invalid admin password.</div>" if error else ""
    return f"<html><head><title>Admin Login</title>{CSS}</head><body><div class='login-box'><h1>Admin Login</h1><p class='subtitle'>Backend access is restricted to admin users only.</p>{message}<form action='/admin-login' method='post'><label><b>Admin Password</b></label><input type='password' name='password' required><br><br><button type='submit'>Access Backend</button></form></div></body></html>"


@application.post("/admin-login")
async def admin_login_submit(password: str = Form("")):
    if password == ADMIN_PASSWORD:
        return RedirectResponse(f"/backend?admin_key={quote(ADMIN_PASSWORD)}", status_code=303)
    return RedirectResponse("/admin-login?error=1", status_code=303)


def test_sharepoint_connection():
    """Return a detailed diagnostic without exposing credentials."""
    diagnostics = {
        "configured": sharepoint_is_configured(),
        "hostname": SHAREPOINT_HOSTNAME,
        "site_path": SHAREPOINT_SITE_PATH,
        "drive_name": SHAREPOINT_DRIVE_NAME,
        "folder_path": SHAREPOINT_FOLDER_PATH,
    }
    if not diagnostics["configured"]:
        missing = [name for name, value in {
            "MS_TENANT_ID": MS_TENANT_ID,
            "MS_CLIENT_ID": MS_CLIENT_ID,
            "MS_CLIENT_SECRET": MS_CLIENT_SECRET,
        }.items() if not value]
        raise RuntimeError("Missing Render environment variables: " + ", ".join(missing))
    token = get_graph_access_token()
    site_id, drive_id, drive_name = resolve_sharepoint_site_and_drive(token)
    ensure_sharepoint_folder(token, drive_id, SHAREPOINT_FOLDER_PATH)
    diagnostics.update({"site_id": site_id, "drive_id": drive_id, "resolved_drive_name": drive_name})
    return diagnostics


@application.get("/admin/sharepoint-test", response_class=HTMLResponse)
async def admin_sharepoint_test(request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)
    try:
        details = test_sharepoint_connection()
        message = (
            "SharePoint connection successful. "
            f"Site: {safe_text(details['site_path'])}; "
            f"Library: {safe_text(details['resolved_drive_name'])}; "
            f"Folder: {safe_text(details['folder_path'])}."
        )
        css_class = "success"
    except Exception as exc:
        message = "SharePoint connection failed: " + str(exc)
        css_class = "error"
    return f"""
    <html><head><title>SharePoint Test</title>{CSS}</head><body><div class='container'>
      <div class='brand-header'><img class='logo' src='data:image/png;base64,{LOGO_BASE64}' alt='Protean Medico Legal'><h1>SharePoint Connection Test</h1></div>
      <div class='{css_class}'>{safe_text(message)}</div>
      <div class='actions'><a class='btn' href='/backend?admin_key={quote(ADMIN_PASSWORD)}'>Back to Backend</a></div>
    </div></body></html>
    """


@application.post("/admin/sharepoint-retry")
async def admin_sharepoint_retry(request: Request, document_id: int = Form(...)):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)
    with db_connection() as conn:
        row = conn.execute(
            """SELECT submission_reference, original_filename, stored_filename, content_type, file_size
               FROM booking_documents WHERE id=?""",
            (document_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Document record not found")
    submission_reference, original_filename, stored_filename, content_type, file_size = row
    file_path = (UPLOAD_DIR / stored_filename).resolve()
    if file_path.parent != UPLOAD_DIR or not file_path.is_file():
        raise HTTPException(status_code=404, detail="Local recovery file not found")
    file_info = {
        "original_filename": original_filename,
        "stored_filename": stored_filename,
        "content_type": content_type or "application/octet-stream",
        "file_size": file_size,
        "path": file_path,
    }
    try:
        sync_files_to_sharepoint([file_info], submission_reference)
        with db_connection() as conn:
            conn.execute(
                """UPDATE booking_documents SET sharepoint_url=?, sharepoint_item_id=?,
                   sharepoint_status='Uploaded', sharepoint_error='' WHERE id=?""",
                (file_info.get("sharepoint_url", ""), file_info.get("sharepoint_item_id", ""), document_id),
            )
    except Exception as exc:
        with db_connection() as conn:
            conn.execute(
                "UPDATE booking_documents SET sharepoint_status='Pending retry', sharepoint_error=? WHERE id=?",
                (str(exc)[:2000], document_id),
            )
    return RedirectResponse(f"/backend?admin_key={quote(ADMIN_PASSWORD)}", status_code=303)


@application.get("/backend", response_class=HTMLResponse)
async def backend(request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)

    headers, rows = fetch_bookings()
    booking_ids = [row[0] for row in rows]
    documents_by_booking = fetch_documents_for_bookings(booking_ids)

    header_html = "".join(f"<th>{safe_text(h)}</th>" for h in headers)
    header_html += "<th>Uploaded Documents</th><th>Action</th>"
    body_html = ""

    for row in rows:
        booking_id = row[0]
        body_html += "<tr>"
        for value in row:
            body_html += f"<td>{safe_text(value)}</td>"

        documents = documents_by_booking.get(booking_id, [])
        if documents:
            links = ""
            for doc in documents:
                status = safe_text(doc.get("sharepoint_status", "Pending"))
                error = safe_text(doc.get("sharepoint_error", ""))
                sp_link = (
                    f" <a class='btn btn-green' target='_blank' href='{safe_text(doc['sharepoint_url'])}'>SharePoint</a>"
                    if doc.get("sharepoint_url") else ""
                )
                retry = ""
                if doc.get("sharepoint_status") != "Uploaded":
                    retry = f"""
                    <form style='display:inline' action='/admin/sharepoint-retry?admin_key={quote(ADMIN_PASSWORD)}' method='post'>
                      <input type='hidden' name='document_id' value='{doc['id']}'>
                      <button type='submit' class='btn-orange'>Retry Upload</button>
                    </form>"""
                error_html = f"<div class='small' title='{error}'>Error: {error[:180]}</div>" if error else ""
                links += (
                    f"<li><a href='/admin/document/{doc['id']}?admin_key={quote(ADMIN_PASSWORD)}'>{safe_text(doc['name'])}</a> "
                    f"<span class='small'>({format_size(doc['size'])}) — SharePoint: {status}</span> "
                    f"{sp_link}{retry}{error_html}</li>"
                )
            body_html += f"<td><ul class='file-list'>{links}</ul></td>"
        else:
            body_html += "<td><span class='small'>No documents</span></td>"

        body_html += f"""
        <td class="nowrap">
            <form action="/delete-booking?admin_key={quote(ADMIN_PASSWORD)}" method="post" onsubmit="return confirm('Delete this booking and its document links?');">
                <input type="hidden" name="booking_id" value="{booking_id}">
                <button type="submit" class="btn-red">Delete</button>
            </form>
        </td></tr>
        """

    teams_status = "Configured" if TEAMS_WEBHOOK_URL else "Not configured"
    return f"""
    <html><head><title>Protean Backend</title>{CSS}</head><body><div class="container">
        <div class="brand-header">
            <img class="logo" src="data:image/png;base64,{LOGO_BASE64}" alt="Protean Medico Legal">
            <h1>Admin Backend</h1>
        </div>
        <div class="nav">
            <b>Protean Booking System <span class="badge">Admin Backend</span></b>
            <div>
                <a class="btn" href="/client">Client Interface</a>
                <a class="btn" href="/admin/documents?admin_key={quote(ADMIN_PASSWORD)}">All Documents</a>
                <a class="btn btn-orange" href="/admin/sharepoint-test?admin_key={quote(ADMIN_PASSWORD)}">Test SharePoint</a>
                <a class="btn" href="/admin/backup?admin_key={quote(ADMIN_PASSWORD)}">Download Backup</a>
                <a class="btn btn-green" href="/export?admin_key={quote(ADMIN_PASSWORD)}">Export CSV</a>
                <a class="btn btn-orange" href="/export-excel?admin_key={quote(ADMIN_PASSWORD)}">Export Excel</a>
            </div>
        </div>
        <div class="notice">Microsoft Teams notification: {teams_status}. Data is stored in <b>{safe_text(DATA_DIR)}</b>; uploaded files in <b>{safe_text(UPLOAD_DIR)}</b>; automatic database backups in <b>{safe_text(BACKUP_DIR)}</b>.</div>
        <div class="table-wrap"><table><tr>{header_html}</tr>{body_html}</table></div>
    </div></body></html>
    """


@application.get("/admin/document/{document_id}")
async def download_document(document_id: int, request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)
    with db_connection() as conn:
        row = conn.execute(
            "SELECT original_filename, stored_filename, content_type FROM booking_documents WHERE id=?",
            (document_id,),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Document record not found")
    original_filename, stored_filename, content_type = row
    file_path = (UPLOAD_DIR / stored_filename).resolve()
    if file_path.parent != UPLOAD_DIR or not file_path.is_file():
        raise HTTPException(status_code=404, detail="Document file not found")
    return FileResponse(
        file_path,
        media_type=content_type or "application/octet-stream",
        filename=original_filename,
    )


@application.post("/delete-booking")
async def delete_booking(request: Request, booking_id: int = Form(...)):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)

    # Soft-delete only: preserve the booking, all fields and every uploaded file.
    archived_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db_connection() as conn:
        conn.execute(
            "UPDATE bookings SET is_archived=1, archived_at=? WHERE id=?",
            (archived_at, booking_id),
        )
        conn.execute(
            "UPDATE booking_documents SET is_archived=1 WHERE booking_id=?",
            (booking_id,),
        )
    backup_database()

    return RedirectResponse(
        f"/backend?admin_key={quote(ADMIN_PASSWORD)}",
        status_code=303,
    )


@application.get("/admin/documents", response_class=HTMLResponse)
async def all_documents(request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)

    rows = fetch_all_documents()
    body_html = ""

    for (
        document_id, booking_id, submission_reference, original_filename,
        content_type, file_size, uploaded_at, claimant_name, law_firm
    ) in rows:
        linked_to = (
            f"{safe_text(claimant_name)} — {safe_text(law_firm)}"
            if booking_id is not None
            else "<span class='small'>Archived: original booking deleted</span>"
        )
        body_html += f"""
        <tr>
            <td>{document_id}</td>
            <td>{safe_text(submission_reference)}</td>
            <td>{linked_to}</td>
            <td>{safe_text(original_filename)}</td>
            <td>{safe_text(content_type)}</td>
            <td>{format_size(file_size)}</td>
            <td>{safe_text(uploaded_at)}</td>
            <td>
                <a class="btn btn-green"
                   href="/admin/document/{document_id}?admin_key={quote(ADMIN_PASSWORD)}">
                    Download
                </a>
            </td>
        </tr>
        """

    if not rows:
        body_html = "<tr><td colspan='8'>No documents have been uploaded.</td></tr>"

    return f"""
    <html><head><title>All Uploaded Documents</title>{CSS}</head>
    <body><div class="container">
        <div class="brand-header">
            <img class="logo" src="data:image/png;base64,{LOGO_BASE64}" alt="Protean Medico Legal">
            <h1>All Uploaded Documents</h1>
            <p class="subtitle">Permanent document archive</p>
        </div>
        <div class="nav">
            <b>Protean Booking System <span class="badge">Admin Backend</span></b>
            <div>
                <a class="btn" href="/backend?admin_key={quote(ADMIN_PASSWORD)}">Back to Backend</a>
                <a class="btn" href="/client">Client Interface</a>
            </div>
        </div>
        <div class="notice">
            Documents remain in this archive even when a booking is deleted.
            Storage location: <b>{safe_text(UPLOAD_DIR)}</b>
        </div>
        <div class="table-wrap">
            <table>
                <tr>
                    <th>ID</th>
                    <th>Submission Reference</th>
                    <th>Linked Booking</th>
                    <th>File Name</th>
                    <th>File Type</th>
                    <th>File Size</th>
                    <th>Uploaded At</th>
                    <th>Action</th>
                </tr>
                {body_html}
            </table>
        </div>
    </div></body></html>
    """


@application.get("/admin/backup")
async def create_manual_backup(request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)
    backup_database()
    latest = max(BACKUP_DIR.glob("protean_bookings_*.db"), key=lambda x: x.stat().st_mtime)
    return FileResponse(latest, media_type="application/octet-stream", filename=latest.name)


@application.get("/export")
async def export_csv(request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)
    headers, rows = fetch_bookings()
    output = EXPORT_DIR / "protean_booking_export.csv"
    with output.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    return FileResponse(output, media_type="text/csv", filename="protean_booking_export.csv")


@application.get("/export-excel")
async def export_excel(request: Request):
    if not is_admin(request):
        return RedirectResponse("/admin-login", status_code=303)
    headers, rows = fetch_bookings()
    output = EXPORT_DIR / "protean_booking_export.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Protean Bookings"
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    wb.save(output)
    return FileResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="protean_booking_export.xlsx",
    )
